"""Sinh bản .XLS (BIFF8, như tệp IPCAS thật) của fixture dư nợ: đọc 20.10.2026.xlsx (make_fixtures.py) và ghi lại nguyên các ô
(ngày giữ định dạng ngày) ⇒ 20.10.2026.XLS. Đáp số dùng chung expected.json. Chạy: PYTHONPATH=<pylib> python3 make_xls_fixture.py (cần openpyxl, xlwt)."""
import os, datetime
import openpyxl, xlwt
here = os.path.dirname(os.path.abspath(__file__))
ws = openpyxl.load_workbook(os.path.join(here, '20.10.2026.xlsx'), read_only=True).worksheets[0]
wb = xlwt.Workbook(encoding='utf-8'); sh = wb.add_sheet('Sheet1'); dfmt = xlwt.easyxf(num_format_str='dd/mm/yyyy')
n = 0
for r, row in enumerate(ws.iter_rows(values_only=True)):
    for c, v in enumerate(row):
        if v is None: continue
        if isinstance(v, datetime.datetime): sh.write(r, c, v, dfmt)
        else: sh.write(r, c, v)
    n += 1
wb.save(os.path.join(here, '20.10.2026.XLS'))
print('OK', n, 'dòng')
